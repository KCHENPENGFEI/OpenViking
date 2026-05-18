import json
from pathlib import Path

import httpx
import pytest
from openpyxl import Workbook, load_workbook


@pytest.fixture
def sample_xlsx(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    # header row (ignored by start-row=2)
    ws.append(["novel", "extra", "system_prompt", "user_prompt", "x", "y", "z", "response"])
    # data row 1
    ws.append(["三体", None, "你是这部小说的专家。", "请总结。", None, None, None, None])
    # data row 2 - missing marker -> per-row error
    ws.append(["活着", None, "你是个文学教授。", "请总结。", None, None, None, None])
    out = tmp_path / "queries.xlsx"
    wb.save(out)
    return out


# Capture the truly original httpx classes at module import time so repeated
# patch_*() calls within one test don't stack patches.
_REAL_HTTPX_ASYNC_CLIENT = httpx.AsyncClient


def _patch_async_client(monkeypatch, transport):
    """Patch the AsyncClient used by AsyncOvClient with a MockTransport."""

    def patched_client(*args, **kwargs):
        kwargs["transport"] = httpx.MockTransport(transport)
        return _REAL_HTTPX_ASYNC_CLIENT(*args, **kwargs)

    monkeypatch.setattr("scripts.novel_eval._common.httpx.AsyncClient", patched_client)


def _argv(**overrides) -> list:
    base = {
        "--user-api-key": "k",
        "--user-id": "test-user",
        "--model-name": "doubao-seed-2-0-pro-260215",
        "--base-url": "http://x",
        "--session-id": "fixed-session",
        "--sleep": "0",
    }
    base.update({k: str(v) for k, v in overrides.items()})
    flat: list = ["chat_eval.py"]
    for k, v in base.items():
        flat.extend([k, v])
    return flat


def test_chat_eval_writes_xlsx_and_judge(sample_xlsx: Path, tmp_path: Path, monkeypatch):
    captured: list[httpx.Request] = []

    def transport(request: httpx.Request) -> httpx.Response:
        captured.append(request)
        return httpx.Response(
            200,
            json={
                "session_id": "s1",
                "message": f"answered {len(captured)}",
                "token_usage": {"prompt_tokens": 100, "completion_tokens": 50, "total_tokens": 150},
                "iteration": 3,
                "tools_used_names": ["openviking_search", "openviking_read"],
            },
        )

    out_xlsx = tmp_path / "out.xlsx"
    out_judge = tmp_path / "out.judge.json"

    import scripts.novel_eval.chat_eval as ce

    monkeypatch.setattr(
        "sys.argv",
        _argv(
            **{
                "--xlsx": str(sample_xlsx),
                "--out-xlsx": str(out_xlsx),
                "--out-judge": str(out_judge),
            }
        ),
    )
    _patch_async_client(monkeypatch, transport)

    rc = ce.main()
    # Row 3 lacks "小说" marker -> [ERROR] -> failures > 0 -> exit code 1
    assert rc == 1

    # Excel: header at row 1 col 8 = model name; row 2 col 8 = response with metrics
    wb = load_workbook(out_xlsx)
    ws = wb.active
    assert ws.cell(row=1, column=8).value == "doubao-seed-2-0-pro-260215"
    row2_cell = str(ws.cell(row=2, column=8).value)
    assert "answered 1" in row2_cell
    assert "[metrics]" in row2_cell
    assert "tokens(p/c/t)=100/50/150" in row2_cell
    assert "iter=3" in row2_cell
    assert "tools=[openviking_search,openviking_read]" in row2_cell
    assert "[ERROR]" in str(ws.cell(row=3, column=8).value)

    # Last row = average row
    last = ws.max_row
    assert ws.cell(row=last, column=1).value == "__AVERAGE__"
    avg_cell = str(ws.cell(row=last, column=8).value)
    assert "[metrics-avg]" in avg_cell
    assert "n=1" in avg_cell
    assert "iter_avg=3.00" in avg_cell
    assert "openviking_search:1" in avg_cell
    assert "openviking_read:1" in avg_cell

    # Judge JSON: bot_response is CLEAN (no [metrics] suffix)
    payload = json.loads(out_judge.read_text(encoding="utf-8"))
    assert payload["summary"]["total"] == 2
    assert payload["records"][0]["bot_response"] == "answered 1"
    assert "[metrics]" not in payload["records"][0]["bot_response"]
    assert "[ERROR]" in payload["records"][1]["bot_response"]

    # RAG hint is always injected now (regardless of model name)
    sent_body = json.loads(captured[0].content)
    assert "openviking_search" in sent_body["message"]
    assert "viking://resources/" in sent_body["message"]
    # user_id MUST be in body so vikingbot doesn't fall back to "anonymous"
    assert sent_body["user_id"] == "test-user"
    # Path / auth still right
    assert captured[0].url.path == "/bot/v1/chat"
    assert captured[0].headers["X-API-Key"] == "k"
    # No --account-id supplied → no scoping query param
    assert captured[0].url.params.get("account_id") is None


def test_rag_hint_is_injected_regardless_of_model(sample_xlsx: Path, tmp_path: Path, monkeypatch):
    sent_messages: list[str] = []

    def transport(request: httpx.Request) -> httpx.Response:
        sent_messages.append(json.loads(request.content)["message"])
        return httpx.Response(200, json={"session_id": "s", "message": "ok"})

    out_xlsx = tmp_path / "out.xlsx"
    out_judge = tmp_path / "out.judge.json"

    import scripts.novel_eval.chat_eval as ce

    monkeypatch.setattr(
        "sys.argv",
        _argv(
            **{
                "--xlsx": str(sample_xlsx),
                "--model-name": "volcengine/doubao-seed-2-0-lite-260428",
                "--out-xlsx": str(out_xlsx),
                "--out-judge": str(out_judge),
            }
        ),
    )
    _patch_async_client(monkeypatch, transport)
    ce.main()

    assert len(sent_messages) == 1  # row 3 fails marker check before /chat
    msg = sent_messages[0]
    assert "openviking_search" in msg
    assert "viking://resources/" in msg
    wb = load_workbook(out_xlsx)
    assert wb.active.cell(row=1, column=8).value == "volcengine/doubao-seed-2-0-lite-260428"


def test_response_col_is_configurable(sample_xlsx: Path, tmp_path: Path, monkeypatch):
    def transport(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json={"session_id": "s", "message": "pro-says-hi"})

    out_xlsx = tmp_path / "out.xlsx"
    out_judge = tmp_path / "out.judge.json"

    import scripts.novel_eval.chat_eval as ce

    monkeypatch.setattr(
        "sys.argv",
        _argv(
            **{
                "--xlsx": str(sample_xlsx),
                "--response-col": "9",
                "--out-xlsx": str(out_xlsx),
                "--out-judge": str(out_judge),
            }
        ),
    )
    _patch_async_client(monkeypatch, transport)
    ce.main()

    wb = load_workbook(out_xlsx)
    ws = wb.active
    assert ws.cell(row=1, column=9).value == "doubao-seed-2-0-pro-260215"
    assert "pro-says-hi" in str(ws.cell(row=2, column=9).value)
    # Col 8 untouched
    assert ws.cell(row=1, column=8).value == "response"
    assert ws.cell(row=2, column=8).value is None


def test_chat_eval_edits_existing_out_xlsx_in_place(sample_xlsx: Path, tmp_path: Path, monkeypatch):
    """A second run with the same --out-xlsx must preserve the first run's column,
    and the second run must skip the __AVERAGE__ row from the first run."""
    out_xlsx = tmp_path / "shared.xlsx"
    out_judge_a = tmp_path / "a.judge.json"
    out_judge_b = tmp_path / "b.judge.json"

    def transport_a(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json={"session_id": "s", "message": "ans-from-A"})

    def transport_b(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json={"session_id": "s", "message": "ans-from-B"})

    import scripts.novel_eval.chat_eval as ce

    # Run A: col 8 with model "model-A"
    monkeypatch.setattr(
        "sys.argv",
        _argv(
            **{
                "--xlsx": str(sample_xlsx),
                "--model-name": "model-A",
                "--response-col": "8",
                "--out-xlsx": str(out_xlsx),
                "--out-judge": str(out_judge_a),
            }
        ),
    )
    _patch_async_client(monkeypatch, transport_a)
    ce.main()

    # Run B: col 9 with model "model-B" on the SAME out_xlsx
    monkeypatch.setattr(
        "sys.argv",
        _argv(
            **{
                "--xlsx": str(sample_xlsx),
                "--model-name": "model-B",
                "--response-col": "9",
                "--out-xlsx": str(out_xlsx),
                "--out-judge": str(out_judge_b),
            }
        ),
    )
    _patch_async_client(monkeypatch, transport_b)
    ce.main()

    wb = load_workbook(out_xlsx)
    ws = wb.active
    # A's column survived
    assert ws.cell(row=1, column=8).value == "model-A"
    assert "ans-from-A" in str(ws.cell(row=2, column=8).value)
    # B's column appeared on the original data rows (not the __AVERAGE__ row!)
    assert ws.cell(row=1, column=9).value == "model-B"
    assert "ans-from-B" in str(ws.cell(row=2, column=9).value)
    # Both runs share the SAME __AVERAGE__ row (placed right after the last
    # data row), each carrying its own metric in its own column.
    avg_rows = [
        r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == "__AVERAGE__"
    ]
    assert len(avg_rows) == 1
    avg_r = avg_rows[0]
    assert "[metrics-avg]" in str(ws.cell(row=avg_r, column=8).value)  # model-A's avg
    assert "[metrics-avg]" in str(ws.cell(row=avg_r, column=9).value)  # model-B's avg


def _make_multi_row_xlsx(tmp_path: Path, n: int = 5) -> Path:
    """Build a small xlsx with `n` independent data rows for session tests."""
    wb = Workbook()
    ws = wb.active
    ws.append(["novel", "extra", "system_prompt", "user_prompt", "x", "y", "z", "response"])
    for i in range(n):
        ws.append([f"novel-{i}", None, "讲这部小说。", f"q{i}", None, None, None, None])
    path = tmp_path / "multi.xlsx"
    wb.save(path)
    return path


def test_concurrent_runs_use_shared_session_opt_in(sample_xlsx: Path, tmp_path: Path, monkeypatch):
    """With --no-session-per-row, all rows must carry the same session_id.

    This used to be the default; --no-session-per-row keeps it available for
    multi-turn coherence experiments.
    """
    seen_sessions: list[str] = []

    def transport(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content)
        seen_sessions.append(body["session_id"])
        return httpx.Response(200, json={"session_id": body["session_id"], "message": "ok"})

    big_xlsx = _make_multi_row_xlsx(tmp_path, n=5)
    out_xlsx = tmp_path / "out.xlsx"
    out_judge = tmp_path / "out.judge.json"

    import scripts.novel_eval.chat_eval as ce

    argv = _argv(
        **{
            "--xlsx": str(big_xlsx),
            "--concurrency": "3",
            "--session-id": "shared-S",
            "--out-xlsx": str(out_xlsx),
            "--out-judge": str(out_judge),
        }
    )
    argv.append("--no-session-per-row")
    monkeypatch.setattr("sys.argv", argv)
    _patch_async_client(monkeypatch, transport)
    rc = ce.main()
    assert rc == 0
    assert len(seen_sessions) == 5
    assert all(s == "shared-S" for s in seen_sessions)


def test_session_per_row_is_default(sample_xlsx: Path, tmp_path: Path, monkeypatch):
    """Default behavior: each row gets a unique session id `{base}-row{N:03d}`.

    This prevents in-context drift where the model, after seeing a few prior
    direct-answer turns in shared session history, stops calling required tools
    on later turns.
    """
    seen_sessions: list[str] = []

    def transport(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content)
        seen_sessions.append(body["session_id"])
        return httpx.Response(200, json={"session_id": body["session_id"], "message": "ok"})

    big_xlsx = _make_multi_row_xlsx(tmp_path, n=5)
    out_xlsx = tmp_path / "out.xlsx"
    out_judge = tmp_path / "out.judge.json"

    import scripts.novel_eval.chat_eval as ce

    monkeypatch.setattr(
        "sys.argv",
        _argv(
            **{
                "--xlsx": str(big_xlsx),
                "--concurrency": "3",
                "--session-id": "base-S",
                "--out-xlsx": str(out_xlsx),
                "--out-judge": str(out_judge),
            }
        ),
    )
    _patch_async_client(monkeypatch, transport)
    rc = ce.main()
    assert rc == 0
    assert len(seen_sessions) == 5
    # Rows are 2..6 in the xlsx (row 1 is header). Each gets its own id.
    assert sorted(seen_sessions) == [
        "base-S-row002",
        "base-S-row003",
        "base-S-row004",
        "base-S-row005",
        "base-S-row006",
    ]
    assert len(set(seen_sessions)) == 5  # all unique


def test_account_id_propagates_as_query_param(sample_xlsx: Path, tmp_path: Path, monkeypatch):
    captured: list[httpx.Request] = []

    def transport(request: httpx.Request) -> httpx.Response:
        captured.append(request)
        return httpx.Response(200, json={"session_id": "s", "message": "ok"})

    out_xlsx = tmp_path / "out.xlsx"
    out_judge = tmp_path / "out.judge.json"

    import scripts.novel_eval.chat_eval as ce

    monkeypatch.setattr(
        "sys.argv",
        _argv(
            **{
                "--xlsx": str(sample_xlsx),
                "--account-id": "qijing_ai_doubao_pro",
                "--user-id": "doubao_pro",
                "--out-xlsx": str(out_xlsx),
                "--out-judge": str(out_judge),
            }
        ),
    )
    _patch_async_client(monkeypatch, transport)
    ce.main()

    sent = captured[0]
    assert sent.url.params.get("account_id") == "qijing_ai_doubao_pro"
    body = json.loads(sent.content)
    assert body["user_id"] == "doubao_pro"


def test_invalid_concurrency_returns_two(sample_xlsx: Path, tmp_path: Path, monkeypatch):
    import scripts.novel_eval.chat_eval as ce

    out_xlsx = tmp_path / "o.xlsx"
    out_judge = tmp_path / "o.judge.json"
    monkeypatch.setattr(
        "sys.argv",
        _argv(
            **{
                "--xlsx": str(sample_xlsx),
                "--concurrency": "0",
                "--out-xlsx": str(out_xlsx),
                "--out-judge": str(out_judge),
            }
        ),
    )
    _patch_async_client(monkeypatch, lambda r: httpx.Response(500))
    assert ce.main() == 2
